from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO

def generate_bid_pack_docx(verdict: str, verified_fields: dict, rule_results: list, profile: dict = None) -> BytesIO:
    doc = Document()

    # Document Title
    title = doc.add_heading("TenderMind AI — Executive Bid Pack", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    company_name = profile.get("company_name", "Apex Engineering & Tech Solutions Pvt Ltd") if profile else "Apex Engineering & Tech Solutions Pvt Ltd"

    p_meta = doc.add_paragraph()
    p_meta.add_run(f"Contractor: {company_name}\n").bold = True
    p_meta.add_run(f"System Recommendation: ").bold = True
    
    run_verdict = p_meta.add_run(f"{verdict}\n")
    run_verdict.bold = True
    if "BID (Qualified)" in verdict:
        run_verdict.font.color.rgb = RGBColor(0x00, 0x80, 0x00)
    elif "NO-BID" in verdict:
        run_verdict.font.color.rgb = RGBColor(0xD9, 0x26, 0x26)
    else:
        run_verdict.font.color.rgb = RGBColor(0xD9, 0x9B, 0x00)

    # 1. Extracted & Verified Requirements
    doc.add_heading("1. Extracted & Verified Procurement Requirements", level=1)
    
    for field_name, data in verified_fields.items():
        clean_name = field_name.replace('_', ' ').title()
        if isinstance(data, dict):
            status = data.get("status", "N/A").upper()
            val = data.get("value", "Not Specified")
            page = data.get("page", 0)
            quote = data.get("quote", "")

            p_field = doc.add_paragraph(style="List Bullet")
            p_field.add_run(f"{clean_name}: ").bold = True
            p_field.add_run(f"{val} ")
            p_field.add_run(f"(Page {page}) [{status}]")
            
            if quote:
                p_quote = doc.add_paragraph(style="Normal")
                p_quote.paragraph_format.left_indent = Inches(0.4)
                p_quote.add_run(f'Source Excerpt: "{quote}"').italic = True

        elif isinstance(data, list):
            doc.add_paragraph(f"{clean_name}:", style="List Bullet")
            for item in data:
                if isinstance(item, dict):
                    status = item.get("status", "N/A").upper()
                    val = item.get("value", "")
                    page = item.get("page", 0)
                    quote = item.get("quote", "")
                    p_sub = doc.add_paragraph(style="Normal")
                    p_sub.paragraph_format.left_indent = Inches(0.4)
                    p_sub.add_run(f"• {val} (Page {page}) [{status}]")
                    if quote:
                        p_sub_quote = doc.add_paragraph(style="Normal")
                        p_sub_quote.paragraph_format.left_indent = Inches(0.6)
                        p_sub_quote.add_run(f'Excerpt: "{quote}"').italic = True

    # 2. Rule Evaluation Details
    doc.add_heading("2. Rule Evaluation & Eligibility Audit", level=1)
    for r in rule_results:
        p_rule = doc.add_paragraph(style="List Bullet")
        status_text = r.get("status", "")
        p_rule.add_run(f"[{status_text}] ").bold = True
        p_rule.add_run(f"{r.get('name')} ({r.get('severity', '').upper()} RULE): ").bold = True
        p_rule.add_run(f"{r.get('reason')}")
        if r.get("evidence"):
            p_ev = doc.add_paragraph(style="Normal")
            p_ev.paragraph_format.left_indent = Inches(0.4)
            p_ev.add_run(f"Grounding Evidence: {r.get('evidence')}").italic = True

    # 3. Draft Compliance Statement
    doc.add_heading("3. Draft Compliance Statement", level=1)
    doc.add_paragraph(
        "We hereby confirm that our firm meets all mandatory statutory, financial, and technical criteria "
        "as stipulated in the tender documentation. Necessary earnest money guarantees, verified PEC credentials, "
        "and active taxpayer certificates are enclosed for official evaluation."
    )

    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio
